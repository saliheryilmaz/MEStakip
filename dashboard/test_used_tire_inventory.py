from datetime import timedelta
from io import BytesIO
from urllib.parse import parse_qs, urlparse

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import CikmaLastik, UserProfile


@override_settings(SECURE_SSL_REDIRECT=False, SESSION_COOKIE_SECURE=False,
                   STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
                             'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class UsedTireInventoryTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username='stock-owner')
        cls.other = User.objects.create_user(username='other-owner')

    def setUp(self):
        self.client.force_login(self.user)
        self.url = reverse('dashboard:cikma_lastikler')

    def tire(self, **kwargs):
        values = dict(user=self.user, musteri_adi='Test', marka='Michelin',
                      ebat='205/55R16', mevsim='yaz', adet=4,
                      durum='depolandi', depo_konumu='Depo A')
        values.update(kwargs)
        return CikmaLastik.objects.create(**values)

    def test_summary_counts_units_not_rows_and_excludes_non_stock(self):
        self.tire()
        self.tire(adet=2, durum='cikti')
        self.tire(adet=90, durum='satildi')
        self.tire(adet=80, durum='imha')
        self.tire(adet=70, user=self.other, depo_konumu='Private depot')
        response = self.client.get(self.url)
        self.assertEqual(response.context['stats']['toplam_adet'], 6)
        self.assertEqual(response.context['stats']['toplam_kayit'], 2)
        self.assertEqual(response.context['depot_rows'][0]['adet_toplam'], 6)
        self.assertNotContains(response, 'Private depot')

    def test_removed_missing_depot_filter_is_ignored_without_losing_stock(self):
        for depot in (None, '', '   '):
            self.tire(depo_konumu=depot, adet=2)
        self.tire()
        response = self.client.get(self.url, {'depo_bos': '1'})
        self.assertEqual(response.context['stats']['toplam_adet'], 10)
        self.assertEqual(len(response.context['depot_rows']), 1)
        self.assertEqual(response.context['depot_totals']['adet_toplam'], 10)
        self.assertNotContains(response, 'Depo eksik')
        self.assertNotContains(response, 'Deposu belirtilmemiş')
        self.assertNotContains(response, 'name="depo_bos"')

    def test_depot_selection_is_exact_and_other_filters_combine(self):
        wanted = self.tire(depo_konumu=' Depo A ')
        self.tire(depo_konumu='Depo A1', adet=8)
        self.tire(marka='Pirelli', adet=9)
        self.tire(mevsim='kis', adet=10)
        self.tire(ebat='225/45R17', adet=11)
        response = self.client.get(self.url, dict(depo_konumu='Depo A', marka='Michelin',
                                                 ebat='2055516', mevsim='yaz'))
        self.assertEqual([t.pk for t in response.context['cikma_lastikler']], [wanted.pk])
        self.assertEqual(response.context['stats']['toplam_adet'], 4)
        self.assertEqual(response.context['depot_totals']['adet_toplam'], 12)

    def test_totals_cover_every_page_and_links_preserve_encoded_filters(self):
        depot = 'Raf A&B / Özel + 1'
        for _ in range(51):
            self.tire(depo_konumu=depot, marka='A&B', adet=2)
        response = self.client.get(self.url, {'marka': 'A&B', 'page': '2', 'depo_konumu': depot})
        self.assertEqual(len(response.context['cikma_lastikler']), 1)
        self.assertEqual(response.context['stats']['toplam_adet'], 102)
        query = parse_qs(urlparse(response.context['depot_rows'][0]['url']).query)
        self.assertEqual(query['marka'], ['A&B'])
        self.assertEqual(query['depo_konumu'], [depot])
        self.assertNotIn('page', query)
        self.assertEqual(parse_qs(response.context['filter_query_string'])['depo_konumu'], [depot])

    def test_export_matches_list_and_includes_depot_and_id(self):
        wanted = self.tire(depo_konumu='Merkez')
        self.tire()
        self.tire(durum='imha', depo_konumu=None)
        response = self.client.get(reverse('dashboard:export_cikma_lastikler_excel'), {'depo_konumu': 'Merkez'})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 5).value, 4)
        self.assertEqual(sheet.cell(2, 7).value, 'Merkez')
        self.assertEqual(sheet.cell(2, 8).value, wanted.pk)

    def test_guest_cannot_see_or_probe_depot_details(self):
        UserProfile.objects.update_or_create(user=self.user, defaults={'role': 'misafir'})
        self.tire(user=self.other, depo_konumu='Hidden depot')
        response = self.client.get(self.url, {'depo_konumu': 'Unknown', 'depo_bos': '1'})
        self.assertEqual(response.context['stats']['toplam_adet'], 4)
        self.assertNotContains(response, 'Hidden depot')
        self.assertNotContains(response, 'id="depotSummary"')
        sheet = load_workbook(BytesIO(self.client.get(
            reverse('dashboard:export_cikma_lastikler_excel')).content)).active
        self.assertEqual(sheet.max_column, 6)

    def test_dates_apply_to_summary_and_invalid_dates_do_not_crash(self):
        old = self.tire(adet=10)
        CikmaLastik.objects.filter(pk=old.pk).update(olusturma_tarihi=timezone.now() - timedelta(days=40))
        self.tire(adet=3)
        response = self.client.get(self.url, {'tarih': 'bu-ay'})
        self.assertEqual(response.context['depot_totals']['adet_toplam'], 3)
        self.assertEqual(self.client.get(self.url, {'baslangic_tarihi': '2026-99-99'}).status_code, 200)

    def test_edit_updates_depot_and_quantity_in_summary(self):
        tire = self.tire(depo_konumu=None)
        response = self.client.post(reverse('dashboard:cikma_lastik_duzenle', args=[tire.pk]),
                                    {'depo_konumu': 'Merkez', 'adet': 3, 'durum': 'depolandi'})
        self.assertEqual(response.status_code, 302)
        response = self.client.get(self.url, {'depo_konumu': 'Merkez'})
        self.assertEqual(response.context['stats']['toplam_adet'], 3)
        self.assertEqual(response.context['depot_rows'][0]['depot_name'], 'Merkez')

    def test_empty_stock(self):
        response = self.client.get(self.url)
        self.assertEqual(response.context['stats']['toplam_adet'], 0)
        self.assertContains(response, 'Bu filtrelerde aktif stok bulunamadı.')

    def test_exact_id_search_with_or_without_hash(self):
        wanted = self.tire()
        self.tire(adet=8)
        for query in (str(wanted.pk), f'#{wanted.pk}', f'  {wanted.pk}  '):
            response = self.client.get(self.url, {'id': query})
            self.assertEqual([t.pk for t in response.context['cikma_lastikler']], [wanted.pk])
            self.assertEqual(response.context['stats']['toplam_adet'], 4)
            self.assertEqual(response.context['depot_totals']['adet_toplam'], 4)
            self.assertEqual(parse_qs(response.context['filter_query_string'])['id'], [query.strip()])
            depot_query = parse_qs(urlparse(response.context['depot_rows'][0]['url']).query)
            self.assertEqual(depot_query['id'], [query.strip()])

    def test_id_search_respects_filters_permissions_and_stock_status(self):
        wanted = self.tire()
        for filters in ({'id': wanted.pk, 'marka': 'Pirelli'},
                        {'id': wanted.pk, 'depo_konumu': 'Other'}):
            self.assertEqual(self.client.get(self.url, filters).context['stats']['toplam_kayit'], 0)
        for tire in (self.tire(user=self.other), self.tire(durum='satildi'), self.tire(durum='imha')):
            self.assertEqual(self.client.get(self.url, {'id': tire.pk}).context['stats']['toplam_kayit'], 0)

    def test_invalid_id_is_empty_and_does_not_crash(self):
        self.tire()
        for query in ('abc', '1.5', '-1', '0', '1 OR 1=1', '9' * 100, '9223372036854775808'):
            response = self.client.get(self.url, {'id': query})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.context['stats']['toplam_kayit'], 0)

    def test_id_export_matches_filtered_record(self):
        wanted = self.tire()
        self.tire(adet=8)
        response = self.client.get(reverse('dashboard:export_cikma_lastikler_excel'), {'id': f'#{wanted.pk}'})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 8).value, wanted.pk)

    def test_new_record_is_immediately_in_depot_total(self):
        response = self.client.post(self.url, {'marka': 'Michelin', 'ebat': '205/55R16',
                                              'mevsim': 'yaz', 'adet': 4,
                                              'durum': 'depolandi', 'depo_konumu': 'Merkez'})
        self.assertEqual(response.status_code, 302)
        response = self.client.get(self.url, {'depo_konumu': 'Merkez'})
        self.assertEqual(response.context['stats']['toplam_adet'], 4)

    def test_depot_export_is_literal_text_not_a_formula(self):
        self.tire(depo_konumu='=1+1')
        response = self.client.get(reverse('dashboard:export_cikma_lastikler_excel'))
        cell = load_workbook(BytesIO(response.content)).active.cell(2, 7)
        self.assertEqual(cell.data_type, 's')
        self.assertEqual(cell.value, '=1+1')
